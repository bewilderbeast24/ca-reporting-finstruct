"""Custom Annexures view — Trade Rec/Pay Ageing, Share Capital, Borrowings.

User enters values bucket-by-bucket; live TB tie-out check shows variance.
"""

from __future__ import annotations
import tkinter as tk
from tkinter import ttk, messagebox
from config import THEME as T
from core.annexures import (
    ANNEXURE_DEFS, build_blank_annexure, load_annexure, save_annexure, AnnexureRow
)
from core.wtb_engine import aggregate_by_code, build_wtb_lines, apply_adjustments
from core.master_db import get_lookup_map
from gui.theme import primary_btn, secondary_btn, label, ribbon


class AnnexuresView(ttk.Frame):
    def __init__(self, parent, db, settings_db, on_proceed: callable = None, on_back: callable = None):
        super().__init__(parent)
        self._db   = db
        self._sdb  = settings_db
        self._on_proceed = on_proceed
        self._on_back = on_back
        self._tolerance = settings_db.get_annexure_tolerance()
        self._totals: dict = {}
        self._current_code: str | None = None
        self._current_annx = None
        self._row_entries: list[tuple[tk.StringVar, tk.StringVar]] = []
        self._compute_totals()
        self._build()

    def _compute_totals(self):
        wtb_rows = self._db.get_wtb()
        raw_rows = self._db.get_raw_tb()
        lines    = build_wtb_lines(wtb_rows, raw_rows)
        totals   = aggregate_by_code(lines)
        adj_rows = self._db.get_adjustments()
        if adj_rows:
            totals = apply_adjustments(totals, adj_rows, get_lookup_map())
        self._totals = totals

    def _get_annexure_status(self, code: str) -> str:
        try:
            a = load_annexure(code, self._db, self._totals, self._tolerance)
            has_values = any(r.cy_value != 0 or r.py_value != 0 for r in a.rows)
            if not has_values and (a.tb_total_cy != 0 or a.tb_total_py != 0):
                return "○"
            return "✔" if a.is_balanced else "⚠"
        except Exception:
            return "○"

    def _refresh_combo_status(self):
        current_idx = self._ann_combo.current()
        annx_labels = [(code, defn["title"]) for code, defn in ANNEXURE_DEFS.items()]
        vals = [f"{self._get_annexure_status(c)} {c} — {t}" for c, t in annx_labels]
        self._ann_combo.configure(values=vals)
        if 0 <= current_idx < len(vals):
            self._ann_combo.current(current_idx)

    def _build(self):
        _, top, nav = ribbon(self)
        # Actions on top-right so Proceed stays visible without scrolling
        primary_btn(nav, "Proceed to FS →", command=self._proceed).pack(side="right", padx=2)
        primary_btn(nav, "💾 Save", command=self._save).pack(side="right", padx=2)
        secondary_btn(nav, "📥 Import", command=self._import_excel).pack(side="right", padx=2)
        secondary_btn(nav, "📋 Template", command=self._export_template).pack(side="right", padx=2)
        label(top, "5b.  Custom Annexures",
              style="Sec.TLabel").pack(side="left")
        label(top, "(Ageing, Share Capital, Borrowings)",
              style="Muted.TLabel").pack(side="left", padx=2)
        secondary_btn(top, "← Back to PPE", command=self._back).pack(side="left", padx=2)
        secondary_btn(top, "🔄 Reload TB total", command=self._reload_blank).pack(side="left", padx=2)

        # Tolerance config
        tol_frame = ttk.Frame(self)
        tol_frame.pack(fill="x", padx=8, pady=(0, 4))
        ttk.Label(tol_frame, text="Tie-out tolerance ₹:").pack(side="left")
        self._tol_var = tk.StringVar(value=f"{self._tolerance:.0f}")
        tol_entry = ttk.Entry(tol_frame, textvariable=self._tol_var, width=10)
        tol_entry.pack(side="left", padx=4)
        secondary_btn(tol_frame, "Apply", command=self._set_tolerance).pack(side="left", padx=4)
        ttk.Label(tol_frame,
                  text="(soft limit; user can override on save if needed)",
                  style="Muted.TLabel").pack(side="left", padx=4)

        # Annexure selector
        sel_frame = ttk.Frame(self)
        sel_frame.pack(fill="x", padx=8, pady=4)
        ttk.Label(sel_frame, text="Annexure:").pack(side="left")
        self._ann_var = tk.StringVar()
        annx_labels  = [(code, defn["title"]) for code, defn in ANNEXURE_DEFS.items()]
        self._ann_combo = ttk.Combobox(
            sel_frame, textvariable=self._ann_var,
            values=[f"{self._get_annexure_status(c)} {c} — {t}" for c, t in annx_labels],
            state="readonly", width=74,
        )
        self._ann_combo.pack(side="left", padx=4)
        self._ann_combo.bind("<<ComboboxSelected>>", self._on_annexure_select)
        self._ann_combo.current(0)

        # Header card — TB totals & variance
        hdr = ttk.Frame(self, relief="ridge", borderwidth=2, padding=8)
        hdr.pack(fill="x", padx=8, pady=4)
        self._tb_cy_var  = tk.StringVar(value="—")
        self._tb_py_var  = tk.StringVar(value="—")
        self._sum_cy_var = tk.StringVar(value="—")
        self._sum_py_var = tk.StringVar(value="—")
        self._var_cy_var = tk.StringVar(value="—")
        self._var_py_var = tk.StringVar(value="—")
        self._status_var = tk.StringVar(value="Select an annexure")

        for col, txt in enumerate(["", "CY (₹)", "PY (₹)"]):
            ttk.Label(hdr, text=txt, font=(T["font"], 9, "bold")).grid(row=0, column=col, padx=8)
        ttk.Label(hdr, text="TB Total (from mapping):").grid(row=1, column=0, sticky="w")
        ttk.Label(hdr, textvariable=self._tb_cy_var, foreground=T["primary"]).grid(row=1, column=1)
        ttk.Label(hdr, textvariable=self._tb_py_var, foreground=T["primary"]).grid(row=1, column=2)
        ttk.Label(hdr, text="Annexure Sum:").grid(row=2, column=0, sticky="w")
        ttk.Label(hdr, textvariable=self._sum_cy_var).grid(row=2, column=1)
        ttk.Label(hdr, textvariable=self._sum_py_var).grid(row=2, column=2)
        ttk.Label(hdr, text="Variance:").grid(row=3, column=0, sticky="w")
        self._var_cy_lbl = ttk.Label(hdr, textvariable=self._var_cy_var, font=(T["font"], 9, "bold"))
        self._var_cy_lbl.grid(row=3, column=1)
        self._var_py_lbl = ttk.Label(hdr, textvariable=self._var_py_var, font=(T["font"], 9, "bold"))
        self._var_py_lbl.grid(row=3, column=2)
        self._status_lbl = ttk.Label(hdr, textvariable=self._status_var,
                                     font=(T["font"], 10, "bold"))
        self._status_lbl.grid(row=4, column=0, columnspan=3, pady=(6, 0))

        # Scrollable rows area
        body = ttk.Frame(self)
        body.pack(fill="both", expand=True, padx=8, pady=4)
        canvas = tk.Canvas(body, bg=T["bg"], highlightthickness=0)
        vsb    = ttk.Scrollbar(body, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        self._rows_frame = ttk.Frame(canvas)
        win = canvas.create_window((0, 0), window=self._rows_frame, anchor="nw")
        self._rows_frame.bind("<Configure>", lambda e: canvas.configure(
            scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(win, width=e.width))

        # Load default first annexure
        self._on_annexure_select()

    def _set_tolerance(self):
        try:
            v = float(self._tol_var.get())
            self._sdb.set_annexure_tolerance(v)
            self._tolerance = v
            if self._current_annx:
                self._current_annx.tolerance = v
                self._recompute_and_render()
        except ValueError:
            messagebox.showerror("Bad value", "Enter a numeric tolerance (e.g. 10)")

    def _on_annexure_select(self, *_):
        sel = self._ann_combo.get()
        if not sel:
            return
        clean = sel
        for s in ("✔ ", "⚠ ", "○ "):
            if clean.startswith(s):
                clean = clean[len(s):]
                break
        code = clean.split(" — ")[0].strip()
        self._current_code = code
        self._current_annx = load_annexure(code, self._db, self._totals, self._tolerance)
        self._render_rows()
        self._recompute_and_render()

    def _reload_blank(self):
        if not self._current_code:
            return
        if not messagebox.askyesno("Reload", "Discard entered values and rebuild from TB?"):
            return
        self._current_annx = build_blank_annexure(self._current_code, self._totals, self._tolerance)
        self._render_rows()
        self._recompute_and_render()

    def _render_rows(self):
        for w in self._rows_frame.winfo_children():
            w.destroy()
        self._row_entries = []
        if not self._current_annx:
            return

        # Header row
        for c, txt in enumerate(["Bucket / Sub-Heading", "CY (₹)", "PY (₹)"]):
            ttk.Label(self._rows_frame, text=txt,
                      font=(T["font"], 9, "bold")).grid(row=0, column=c, padx=6, pady=(0, 4),
                                                          sticky="w" if c == 0 else "e")
        for ri, row in enumerate(self._current_annx.rows, start=1):
            ttk.Label(self._rows_frame, text=row.label, anchor="w", wraplength=420).grid(
                row=ri, column=0, padx=6, pady=1, sticky="w")
            cy_v = tk.StringVar(value=f"{row.cy_value:.2f}" if row.cy_value else "0")
            py_v = tk.StringVar(value=f"{row.py_value:.2f}" if row.py_value else "0")
            cy_e = ttk.Entry(self._rows_frame, textvariable=cy_v, width=18, justify="right")
            py_e = ttk.Entry(self._rows_frame, textvariable=py_v, width=18, justify="right")
            cy_e.grid(row=ri, column=1, padx=4, pady=1, sticky="e")
            py_e.grid(row=ri, column=2, padx=4, pady=1, sticky="e")
            cy_v.trace_add("write", lambda *_, i=ri-1: self._on_value_change(i))
            py_v.trace_add("write", lambda *_, i=ri-1: self._on_value_change(i))
            self._row_entries.append((cy_v, py_v))

    def _on_value_change(self, idx: int):
        if not self._current_annx or idx >= len(self._current_annx.rows):
            return
        try:
            cy_str = self._row_entries[idx][0].get().replace(",", "").strip()
            py_str = self._row_entries[idx][1].get().replace(",", "").strip()
            self._current_annx.rows[idx].cy_value = float(cy_str) if cy_str else 0.0
            self._current_annx.rows[idx].py_value = float(py_str) if py_str else 0.0
        except ValueError:
            return
        self._recompute_and_render()

    def _recompute_and_render(self):
        if not self._current_annx:
            return
        self._current_annx.recompute()
        a = self._current_annx
        self._tb_cy_var.set(f"{a.tb_total_cy:,.2f}")
        self._tb_py_var.set(f"{a.tb_total_py:,.2f}")
        sum_cy = sum(r.cy_value for r in a.rows)
        sum_py = sum(r.py_value for r in a.rows)
        self._sum_cy_var.set(f"{sum_cy:,.2f}")
        self._sum_py_var.set(f"{sum_py:,.2f}")
        self._var_cy_var.set(f"{a.variance_cy:+,.2f}")
        self._var_py_var.set(f"{a.variance_py:+,.2f}")
        if a.is_balanced:
            self._status_var.set(f"✅  Tied out to TB (within ₹{a.tolerance:.0f} tolerance)")
            self._status_lbl.configure(foreground=T["success"])
            self._var_cy_lbl.configure(foreground=T["success"])
            self._var_py_lbl.configure(foreground=T["success"])
        else:
            self._status_var.set(
                f"⚠  Annexure does NOT tie out — variance exceeds ₹{a.tolerance:.0f}"
            )
            self._status_lbl.configure(foreground=T["error"])
            self._var_cy_lbl.configure(foreground=T["error"])
            self._var_py_lbl.configure(foreground=T["error"])

    def _save(self):
        if not self._current_annx:
            return
        a = self._current_annx
        if not a.is_balanced:
            msg = (f"Annexure does not tie out to TB total.\n"
                   f"CY variance: ₹{a.variance_cy:+,.2f}\n"
                   f"PY variance: ₹{a.variance_py:+,.2f}\n\n"
                   f"Save anyway? (variance will be logged)")
            if not messagebox.askyesno("Variance", msg):
                return
            self._db.log("ANNEXURE_VARIANCE",
                         f"{a.code}: CY variance ₹{a.variance_cy:.2f}")
        save_annexure(a, self._db)
        self._db.log("ANNEXURE_SAVED", f"{a.code}: {len(a.rows)} rows")
        self._refresh_combo_status()
        messagebox.showinfo("Saved", f"✅ {a.title}\nsaved successfully.")

    def _proceed(self):
        if self._on_proceed:
            self._on_proceed()

    def _back(self):
        if self._on_back:
            self._on_back()

    def _import_excel(self):
        if not self._current_annx:
            return
        from tkinter import filedialog
        path = filedialog.askopenfilename(
            title=f"Import {self._current_code} from Excel",
            filetypes=[("Excel Workbook", "*.xlsx"), ("All Files", "*.*")],
        )
        if not path:
            return
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            label_map = {r.label.strip().lower(): r for r in self._current_annx.rows}
            matched = 0
            for row in ws.iter_rows(values_only=True):
                if not row or row[0] is None:
                    continue
                key = str(row[0]).strip().lower()
                if key in label_map:
                    target = label_map[key]
                    try:
                        target.cy_value = float(row[1]) if len(row) > 1 and row[1] is not None else 0.0
                    except (ValueError, TypeError):
                        pass
                    try:
                        target.py_value = float(row[2]) if len(row) > 2 and row[2] is not None else 0.0
                    except (ValueError, TypeError):
                        pass
                    matched += 1
            if matched == 0:
                # Attempt positional import for rows after header
                data_rows = []
                for row in ws.iter_rows(values_only=True):
                    if not row or row[0] is None:
                        continue
                    if len(row) >= 2 and any(isinstance(c, (int, float)) for c in row[1:3]):
                        data_rows.append(row)
                for i, row in enumerate(data_rows):
                    if i < len(self._current_annx.rows):
                        target = self._current_annx.rows[i]
                        try:
                            target.cy_value = float(row[1]) if len(row) > 1 and row[1] is not None else 0.0
                        except (ValueError, TypeError):
                            pass
                        try:
                            target.py_value = float(row[2]) if len(row) > 2 and row[2] is not None else 0.0
                        except (ValueError, TypeError):
                            pass
                        matched += 1

            self._render_rows()
            self._recompute_and_render()
            self._refresh_combo_status()
            messagebox.showinfo("Import successful", f"Imported {matched} values from Excel.")
        except Exception as e:
            messagebox.showerror("Import failed", f"Could not import Excel file:\n{e}")

    def _export_template(self):
        if not self._current_annx:
            return
        from tkinter import filedialog
        from pathlib import Path
        path = filedialog.asksaveasfilename(
            title=f"Export {self._current_code} template",
            defaultextension=".xlsx",
            initialfile=f"{self._current_code}_template.xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = self._current_code[:31]
            ws.append([self._current_annx.title])
            ws.append([])
            ws.append([f"TB Total CY:  ₹{self._current_annx.tb_total_cy:,.2f}",
                       f"TB Total PY:  ₹{self._current_annx.tb_total_py:,.2f}"])
            ws.append([])
            ws.append(["Bucket / Sub-Heading", "CY (₹)", "PY (₹)"])
            for row in self._current_annx.rows:
                ws.append([row.label, row.cy_value, row.py_value])
            wb.save(path)
            messagebox.showinfo("Exported", f"✅ Template saved to:\n{path}")
        except Exception as e:
            messagebox.showerror("Export failed", str(e))
